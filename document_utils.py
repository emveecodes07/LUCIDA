"""
document_utils.py
Parses uploaded course material (pdf, pptx, docx) and builds prompts that turn
it into a full local "study pack": a revision summary, a plain-language
overview, section pointers, key terms, flashcards, and a short quiz.

Speed note: the study pack is built from ONE model call instead of several.
Every extra call re-sends and re-processes the full source text, which is the
expensive part — a single combined JSON response cuts that cost roughly in
half-to-a-third versus separate summary/quiz/flashcard calls, with no loss of
depth in any individual section.
"""

import hashlib
import json

from pypdf import PdfReader
from pptx import Presentation
from docx import Document

MAX_CHARS = 8000  # keeps prompts within a safe context window for small local models


def extract_text(file) -> str:
    ext = file.name.split(".")[-1].lower()
    text = ""
    try:
        if ext == "pdf":
            reader = PdfReader(file)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        elif ext in ("ppt", "pptx"):
            prs = Presentation(file)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext in ("doc", "docx"):
            doc = Document(file)
            for p in doc.paragraphs:
                text += p.text + "\n"
        else:
            raise ValueError(f"Unsupported file type: .{ext}")
    except Exception as e:
        raise ValueError(f"Couldn't read this {ext.upper()} file: {e}")

    text = text.strip()
    if not text:
        raise ValueError("No extractable text found — is this a scanned/image-only file?")
    return text[:MAX_CHARS]


def file_fingerprint(file) -> str:
    """Cheap hash so we skip regenerating a study pack for a file we've already processed."""
    file.seek(0)
    digest = hashlib.md5(file.read()).hexdigest()
    file.seek(0)
    return digest


def _strip_json_fences(raw: str) -> str:
    cleaned = raw.strip()
    for fence in ("```json", "```"):
        cleaned = cleaned.replace(fence, "")
    return cleaned.strip()


# ---------------------------------------------------------------------------
# Combined study pack: summary + overview + section pointers + keywords +
# flashcards + quiz, all from a single model call.
# ---------------------------------------------------------------------------

def build_study_pack_prompt(text: str, n_quiz: int = 5, n_flashcards: int = 10) -> str:
    return (
        "You are building a complete, exam-ready study pack from the material below. Base every "
        "field ONLY on the material — never invent facts, numbers, or terms that aren't in it. If "
        "the material is too thin for a field (e.g. no real misconceptions to flag, or it's a "
        "short/introductory doc with no real prerequisites), return a shorter list for that field "
        "rather than padding it with filler.\n\n"
        "Respond with ONLY a raw JSON object, no markdown fences, no prose outside the JSON, "
        "matching EXACTLY this shape:\n"
        "{\n"
        '  "overview": "2-3 sentence plain-language overview of what this material covers and why it matters",\n'
        '  "tldr": "ONE sentence — if the student only remembers one thing before a last-minute glance, this is it",\n'
        '  "difficulty": "one word: Beginner, Intermediate, or Advanced — your best judgment of the material\'s level",\n'
        '  "est_study_minutes": integer — a realistic minutes estimate to read + review this pack once,\n'
        '  "prerequisites": ["background knowledge/terms this material assumes you already know — skip entirely if it\'s self-contained/introductory"],\n'
        '  "summary": ["6-9 high-yield bullets — concepts, definitions, facts likely to be tested, no filler"],\n'
        '  "section_pointers": ["one line per major section/topic: name — what to focus on when reviewing it"],\n'
        '  "keywords": ["10-16 major keywords or named concepts, most important first"],\n'
        '  "analogies": [{"concept": "the trickiest concept in the material", "analogy": "a short, vivid everyday comparison that makes it click"}] — 2-4 of these, only for genuinely hard concepts, skip if everything is already simple,\n'
        '  "mnemonics": [{"term": "a list, sequence, or hard-to-recall fact from the material", "device": "a short memory trick — acronym, rhyme, or vivid mental image"}] — only where a real mnemonic fits naturally, skip if nothing in the material lends itself to one,\n'
        '  "common_mistakes": ["3-5 mistakes or misconceptions students commonly make with this material — skip if the material gives no basis for this"],\n'
        '  "real_world_applications": ["2-4 short bullets on where/how this shows up outside the classroom or exam — skip if genuinely purely theoretical with no applied angle"],\n'
        f'  "flashcards": [{{"front": "question or term", "back": "concise answer"}}] — exactly {n_flashcards} cards, '
        "covering different parts of the material, ordered roughly easiest to hardest,\n"
        '  "practice_questions": ["3 short-answer (not multiple-choice) questions that require explaining a concept in the student\'s own words, for active-recall self-testing"],\n'
        f'  "quiz": [{{"q": "question text", "options": ["A) ...", "B) ...", "C) ...", "D) ..."], "answer": "A"}}] — exactly {n_quiz} questions, spanning different sections rather than clustering on one topic,\n'
        '  "study_plan": [{"day": 1, "focus": "what to review this session, referencing real section names", "minutes": integer}] — split est_study_minutes into realistic sessions (use just one entry, day 1, if the material is short enough for a single sitting; use 2-4 entries for denser material, spacing the hardest sections earliest so there\'s time to revisit them)\n'
        "}\n\nMATERIAL:\n" + text
    )


def parse_study_pack_json(raw: str) -> dict:
    """Robustly parses the study pack JSON, defaulting any missing/malformed field to empty."""
    cleaned = _strip_json_fences(raw)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        return {}
    if not isinstance(data, dict):
        return {}

    def _str(v):
        return v if isinstance(v, str) else ""

    def _list_of_str(v):
        return [x for x in v if isinstance(x, str)] if isinstance(v, list) else []

    def _flashcards(v):
        if not isinstance(v, list):
            return []
        return [c for c in v if isinstance(c, dict) and "front" in c and "back" in c]

    def _quiz(v):
        if not isinstance(v, list):
            return []
        return [q for q in v if isinstance(q, dict) and "q" in q and "options" in q and "answer" in q]

    def _keyed_list(v, keys):
        if not isinstance(v, list):
            return []
        return [d for d in v if isinstance(d, dict) and all(k in d for k in keys)]

    def _study_plan(v):
        if not isinstance(v, list):
            return []
        plan = []
        for d in v:
            if not isinstance(d, dict) or "focus" not in d:
                continue
            plan.append({
                "day": _int(d.get("day")) or (len(plan) + 1),
                "focus": _str(d.get("focus")),
                "minutes": _int(d.get("minutes")) or 0,
            })
        return plan

    def _int(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return None

    return {
        "overview": _str(data.get("overview")),
        "tldr": _str(data.get("tldr")),
        "difficulty": _str(data.get("difficulty")),
        "est_study_minutes": _int(data.get("est_study_minutes")),
        "prerequisites": _list_of_str(data.get("prerequisites")),
        "summary": _list_of_str(data.get("summary")),
        "section_pointers": _list_of_str(data.get("section_pointers")),
        "keywords": _list_of_str(data.get("keywords")),
        "analogies": _keyed_list(data.get("analogies"), ("concept", "analogy")),
        "mnemonics": _keyed_list(data.get("mnemonics"), ("term", "device")),
        "common_mistakes": _list_of_str(data.get("common_mistakes")),
        "real_world_applications": _list_of_str(data.get("real_world_applications")),
        "flashcards": _flashcards(data.get("flashcards")),
        "practice_questions": _list_of_str(data.get("practice_questions")),
        "quiz": _quiz(data.get("quiz")),
        "study_plan": _study_plan(data.get("study_plan")),
    }


# ---------------------------------------------------------------------------
# Generic (no-upload) quiz — built from a topic the user types in, using the
# model's own knowledge rather than an uploaded file.
# ---------------------------------------------------------------------------

def build_generic_quiz_prompt(topic: str, n_questions: int = 4) -> str:
    return (
        f"Generate exactly {n_questions} multiple-choice quiz questions to help someone study "
        f'the topic: "{topic}". Use your own general knowledge — no source material is provided. '
        "Respond with ONLY a raw JSON array, no markdown, no prose, matching exactly this shape: "
        '[{"q": "question text", "options": ["A) ...", "B) ...", "C) ...", "D) ..."], "answer": "A"}]'
    )


def study_pack_to_markdown(pack: dict, title: str = "Study Pack") -> str:
    """Renders the full study pack as a single Markdown doc for download —
    lets students keep/print/share it outside the app."""
    lines = [f"# {title}", ""]
    if pack.get("tldr"):
        lines += [f"> **TL;DR:** {pack['tldr']}", ""]
    if pack.get("overview"):
        lines += ["## Overview", pack["overview"], ""]
    badges = []
    if pack.get("difficulty"):
        badges.append(f"Difficulty: {pack['difficulty']}")
    if pack.get("est_study_minutes"):
        badges.append(f"Est. time: ~{pack['est_study_minutes']} min")
    if badges:
        lines += [" · ".join(badges), ""]
    if pack.get("prerequisites"):
        lines += ["## Prerequisites"] + [f"- {p}" for p in pack["prerequisites"]] + [""]
    if pack.get("summary"):
        lines += ["## High-yield summary"] + [f"- {b}" for b in pack["summary"]] + [""]
    if pack.get("section_pointers"):
        lines += ["## Section pointers"] + [f"- {p}" for p in pack["section_pointers"]] + [""]
    if pack.get("keywords"):
        lines += ["## Key terms", ", ".join(pack["keywords"]), ""]
    if pack.get("analogies"):
        lines += ["## Analogies"]
        lines += [f"- **{a['concept']}** — {a['analogy']}" for a in pack["analogies"]]
        lines += [""]
    if pack.get("mnemonics"):
        lines += ["## Memory tricks"]
        lines += [f"- **{m['term']}** — {m['device']}" for m in pack["mnemonics"]]
        lines += [""]
    if pack.get("common_mistakes"):
        lines += ["## Common mistakes"] + [f"- {m}" for m in pack["common_mistakes"]] + [""]
    if pack.get("real_world_applications"):
        lines += ["## Real-world applications"] + [f"- {a}" for a in pack["real_world_applications"]] + [""]
    if pack.get("study_plan"):
        lines += ["## Study plan"]
        for d in pack["study_plan"]:
            mins = f" ({d['minutes']} min)" if d.get("minutes") else ""
            lines += [f"- **Day {d['day']}{mins}:** {d['focus']}"]
        lines += [""]
    if pack.get("practice_questions"):
        lines += ["## Practice questions (self-testing)"]
        lines += [f"{i}. {q}" for i, q in enumerate(pack["practice_questions"], 1)]
        lines += [""]
    if pack.get("flashcards"):
        lines += ["## Flashcards"]
        for i, c in enumerate(pack["flashcards"], 1):
            lines += [f"{i}. **Q:** {c['front']}  \n   **A:** {c['back']}"]
        lines += [""]
    if pack.get("quiz"):
        lines += ["## Quiz"]
        for i, q in enumerate(pack["quiz"], 1):
            lines += [f"{i}. {q['q']}"]
            lines += [f"   - {opt}" for opt in q.get("options", [])]
            lines += [f"   - **Answer:** {q['answer']}", ""]
    return "\n".join(lines).strip() + "\n"


def parse_quiz_json(raw: str) -> list:
    cleaned = _strip_json_fences(raw)
    try:
        data = json.loads(cleaned)
        if isinstance(data, list) and all(
            isinstance(q, dict) and "q" in q and "options" in q and "answer" in q for q in data
        ):
            return data
    except json.JSONDecodeError:
        pass
    return []
